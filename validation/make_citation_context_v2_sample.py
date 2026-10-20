"""
validation/make_citation_context_v2_sample.py
Draws a stratified, seeded sample of citation_context_v2 rows for hand
labelling and writes validation/citation_context_v2_validation.xlsx.

    python -m validation.make_citation_context_v2_sample

Validates (1) fuzzy / anchor matching of the LLM citation_sentence, oversampled
near the score cutoff, and (2) the core-sentence step added in v2. Citations
already in citation_context_validation.xlsx are skipped, and one citation is
drawn per window_key so a labeller never sees the same shared span twice.
Each row carries a window of raw markdown: the matched span in « », the
chosen core in ⟦ ⟧.
"""
import random
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from loaders.citation_context import MD_ROOT

SEED   = 20261013
DB     = Path(__file__).parent.parent / "dockets.db"
OUT    = Path(__file__).parent / "citation_context_v2_validation.xlsx"
V1_SET = Path(__file__).parent / "citation_context_validation.xlsx"
WINDOW = 500
MAX_CELL = 32000

# (stratum, n, WHERE clause on v = citation_context_v2). Drawn in order, so the
# fuzzy strata take priority; the core strata are then filled from what is left.
STRATA = [
    ("fuzzy_85_90",       30, "v.match_method = 'fuzzy' AND v.match_score < 90"),
    ("fuzzy_90_95",       25, "v.match_method = 'fuzzy' AND v.match_score >= 90 AND v.match_score < 95"),
    ("fuzzy_95plus",      20, "v.match_method = 'fuzzy' AND v.match_score >= 95"),
    ("anchor",            25, "v.match_method = 'anchor'"),
    ("marker",            30, "v.core_method = 'marker'"),
    ("span_no_marker",    20, "v.core_method = 'span_no_marker'"),
    ("span_multi_marker", 15, "v.core_method = 'span_multi_marker'"),
    ("single",            10, "v.core_method = 'single' AND v.unit_type = 'body'"),
    ("short_window",      10, "v.core_method IN ('single', 'marker') AND v.n_window_sents < 3"),
]

COLUMNS = [  # (header, key, width, is_label)
    ("val_id", "val_id", 7, False), ("stratum", "stratum", 15, False),
    ("match_method", "match_method", 12, False), ("match_score", "match_score", 8, False),
    ("citation_sentence (LLM)", "citation_sentence", 45, False),
    ("core_method", "core_method", 15, False), ("n_span_sents", "n_span_sents", 7, False),
    ("n_marker_sents", "n_marker_sents", 7, False), ("n_window_sents", "n_window_sents", 7, False),
    ("citation footnote / reference (which source this row is)", "source_hint", 40, False),
    ("prev_sentence", "prev_sentence", 35, False), ("core_sentence", "core_sentence", 45, False),
    ("next_sentence", "next_sentence", 35, False),
    ("prev_boundary", "prev_boundary", 11, False), ("next_boundary", "next_boundary", 11, False),
    ("source_window (LLM span in « », core in ⟦ ⟧)", "window", 70, False),
    ("match_ok", None, 10, True), ("core_ok", None, 10, True),
    ("prev_correct", None, 10, True), ("next_correct", None, 10, True), ("notes", None, 30, True),
    ("citation_id", "citation_id", 38, False), ("comment_id", "comment_id", 22, False),
    ("source_path", "source_path", 40, False),
]
LABELS = {"match_ok": '"Y,Partial,N"', "core_ok": '"Y,N"', "prev_correct": '"Y,N,NA"', "next_correct": '"Y,N,NA"'}


def _v1_ids() -> set[str]:
    if not V1_SET.exists():
        return set()
    ws = load_workbook(V1_SET, read_only=True)["Validation"]
    hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
    i = hdr.index("citation_id")
    return {r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i]}


def _sample(con: sqlite3.Connection) -> list[dict]:
    rng, taken_ids, taken_windows, rows = random.Random(SEED), _v1_ids(), set(), []
    for stratum, n, where in STRATA:
        pool, seen = [], set()
        for r in con.execute(
                "SELECT v.*, x.char_start AS span_start, x.char_end AS span_end, "
                "c.citation_sentence, c.citation_footnote_text, c.citation_reference_text, c.citation_link "
                "FROM citation_context_v2 v JOIN citation_context x USING(citation_id) "
                f"JOIN comment_citations c USING(citation_id) WHERE {where} ORDER BY v.citation_id"):
            if r["window_key"] in taken_windows or r["window_key"] in seen or r["citation_id"] in taken_ids:
                continue
            seen.add(r["window_key"])
            pool.append(r)
        for r in rng.sample(pool, min(n, len(pool))):
            taken_ids.add(r["citation_id"])
            taken_windows.add(r["window_key"])
            rows.append({**dict(r), "stratum": stratum})
    return rows


def _source_hint(r: dict) -> str:
    parts = [(lbl, r[k]) for lbl, k in (("footnote", "citation_footnote_text"),
                                        ("reference", "citation_reference_text"),
                                        ("link", "citation_link")) if r[k]]
    return "\n".join(f"[{lbl}] {v[:400]}" for lbl, v in parts) or "(no footnote/reference text)"


def _window(raw: str, r: dict) -> str:
    s0, s1, c0, c1 = r["span_start"], r["span_end"], r["core_char_start"], r["core_char_end"]
    lo = max(0, min(s0, r["prev_char_start"] or s0) - WINDOW)
    hi = min(len(raw), max(s1, r["next_char_end"] or s1) + WINDOW)
    marks = sorted([(s0, 0, "«"), (c0, 1, "⟦"), (c1, 2, "⟧"), (s1, 3, "»")])
    out, pos = [], lo
    for p, _, m in marks:
        p = min(max(p, lo), hi)
        out += [raw[pos:p], m]
        pos = p
    out.append(raw[pos:hi])
    return (("…" if lo else "") + "".join(out) + ("…" if hi < len(raw) else ""))[:MAX_CELL]


def _write(rows: list[dict]) -> None:
    wb = Workbook()
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill, label_fill = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("Citation context v2 validation set (fuzzy matching + core-sentence selection)", bold),
        (f"{len(rows)} rows from citation_context_v2, one per distinct window (seed {SEED}). The fuzzy_* and "
         "anchor strata check fuzzy matching, oversampling scores near the 85 cutoff; the rest check core "
         "selection. Rows already in citation_context_validation.xlsx are excluded. "
         "Regenerate with: python -m validation.make_citation_context_v2_sample", font),
        ("", font),
        ("What v2 does: when the LLM citation_sentence covers several sentences, v2 keeps only the sentence "
         "carrying the citation marker (core_method = marker) and takes ±1 sentence around it. If no sentence "
         "or several sentences carry markers, the whole span is the core and no neighbours are added "
         "(span_no_marker / span_multi_marker).", font),
        ("Matching: the LLM's citation_sentence is located in the markdown by exact match, else by fuzzy "
         "alignment (match_method fuzzy, score >= 85) or by fuzzy head/tail anchors for long sentences "
         "(match_method anchor, score >= 90).", font),
        ("Compare 'citation_sentence (LLM)' with the « » span in source_window. Use the 'citation footnote / "
         "reference' column to tell which source this row is about.", font),
        ("", font),
        ("Edit only the yellow columns on the 'Validation' tab.", bold),
        ("match_ok — is the « » span the same text as citation_sentence (LLM)?", font),
        ("    Y = same sentence(s) · Partial = right place but cut short or runs over (e.g. stops at a page break "
         "before the marker, or pulls in footnote text) · N = wrong place (similar wording elsewhere, a heading, "
         "a repeated boilerplate passage)", font),
        ("core_ok — depends on core_method:", font),
        ("    single / marker: Y = ⟦core⟧ is the sentence that cites this source · N = wrong sentence "
         "(name the right one in notes)", font),
        ("    span_no_marker: Y = no single sentence in the span carries a marker for this source, so keeping the "
         "span is right · N = a marker was there and missed (copy it into notes)", font),
        ("    span_multi_marker: Y = two or more span sentences really carry citation markers · "
         "N = a detected 'marker' is an ordinary number (e.g. 'Fig. 3', 'No. 5')", font),
        ("prev_correct / next_correct — single/marker rows only: Y = the true neighbouring sentence, or correctly "
         "NULL at a heading/table/edge · N = wrong · NA = span rows (no neighbours by design)", font),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ins.cell(i, 1, t).font = f
    ins.column_dimensions["A"].width = 160

    ws = wb.create_sheet("Validation")
    for j, (h, _, w, is_label) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, h)
        c.font, c.alignment, c.border = bold, wrap, thin
        c.fill = label_fill if is_label else head_fill
        ws.column_dimensions[c.column_letter].width = w
    for i, r in enumerate(rows, 2):
        r["val_id"] = i - 1
        for j, (h, k, _, is_label) in enumerate(COLUMNS, 1):
            v = r.get(k) if k else None
            if h in ("prev_correct", "next_correct") and r["core_method"].startswith("span"):
                v = "NA"
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap, thin
            if is_label:
                c.fill = label_fill
    last = len(rows) + 1
    headers = [h for h, *_ in COLUMNS]
    for name, choices in LABELS.items():
        letter = ws.cell(1, headers.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last}"

    sm = wb.create_sheet("Summary")
    col = {h: ws.cell(1, headers.index(h) + 1).column_letter for h in headers}
    rng = lambda h: f"Validation!${col[h]}$2:${col[h]}${last}"   # noqa: E731
    S = rng("stratum")
    hdr = ["stratum", "n", "labelled", "match Y rate", "core Y rate", "prev Y rate", "next Y rate"]
    for j, h in enumerate(hdr, 1):
        c = sm.cell(1, j, h)
        c.font, c.fill = bold, head_fill
        sm.column_dimensions[c.column_letter].width = 16
    rate = lambda crit, h: (f'=IFERROR(COUNTIFS({S},{crit},{rng(h)},"Y")/(COUNTIFS({S},{crit},{rng(h)},"<>")'  # noqa: E731
                            f'-COUNTIFS({S},{crit},{rng(h)},"NA")),"")')
    for i, s in enumerate([s for s, _, _ in STRATA] + ["ALL"], 2):
        crit = '"*"' if s == "ALL" else f"$A{i}"
        sm.cell(i, 1, s).font = bold if s == "ALL" else font
        sm.cell(i, 2, f"=COUNTIF({S},{crit})")
        sm.cell(i, 3, f'=COUNTIFS({S},{crit},{rng("match_ok")},"<>")')
        for j, h in enumerate(("match_ok", "core_ok", "prev_correct", "next_correct"), 4):
            c = sm.cell(i, j, rate(crit, h))
            c.number_format = "0.0%"
    sm.cell(len(STRATA) + 4, 1, "Rates exclude blank and NA labels; Partial counts as not Y.").font = \
        Font(name="Arial", size=9, italic=True)

    wb.save(OUT)


def main() -> None:
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    rows = _sample(con)
    raws: dict[str, str] = {}
    for r in rows:
        p = r["source_path"]
        if p not in raws:
            raws[p] = (MD_ROOT / p).read_text(encoding="utf-8", errors="ignore")
        r["window"] = _window(raws[p], r)
        r["source_hint"] = _source_hint(r)
    _write(rows)
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["stratum"]] = counts.get(r["stratum"], 0) + 1
    print(f"wrote {len(rows)} rows to {OUT}")
    print("  " + "  ".join(f"{k}={v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
