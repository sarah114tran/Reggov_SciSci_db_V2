"""
validation/make_citation_context_v3_sample.py
Draws a stratified, seeded sample of citation_context_v2 rows (after the
validation-set-2 fixes) for hand labelling and writes
validation/citation_context_v3_validation.xlsx.

    python -m validation.make_citation_context_v3_sample

Strata target the new behaviour: core chosen by footnote number (in place
and relocated), page-break repairs, link / span_last fallbacks, low-score
fuzzy matches and footnote/table matches. Citations already in validation
sets 1 and 2 are skipped, and one citation is drawn per window_key. Each row
carries a window of raw markdown: the matched LLM span in « », the chosen
core in ⟦ ⟧ (they can be apart when the core was relocated).
"""
import random
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from loaders.citation_context import MD_ROOT

SEED   = 20261014
DB     = Path(__file__).parent.parent / "dockets.db"
OUT    = Path(__file__).parent / "citation_context_v3_validation.xlsx"
EARLIER = [Path(__file__).parent / "citation_context_validation.xlsx",
           Path(__file__).parent / "citation_context_v2_validation.xlsx"]
WINDOW = 400
MAX_CELL = 32000

# (stratum, n, WHERE clause on v = citation_context_v2); drawn in order
STRATA = [
    ("relocated",        25, "v.relocated = 1"),
    ("footnote_number",  20, "v.core_method = 'footnote_number' AND v.relocated = 0"),
    ("repaired",         20, "v.repaired = 1"),
    ("span_last",        20, "v.core_method = 'span_last'"),
    ("link",             15, "v.core_method = 'link'"),
    ("fuzzy_85_90",      20, "v.match_method = 'fuzzy' AND v.match_score < 90"),
    ("marker",           10, "v.core_method = 'marker'"),
    ("footnote_table",   10, "v.unit_type IN ('footnote', 'table')"),
]

COLUMNS = [  # (header, key, width, is_label)
    ("val_id", "val_id", 7, False), ("stratum", "stratum", 14, False),
    ("match_method", "match_method", 11, False), ("match_score", "match_score", 7, False),
    ("core_method", "core_method", 14, False), ("citation_number", "citation_number", 8, False),
    ("number_source", "number_source", 12, False), ("relocated", "relocated", 8, False),
    ("repaired", "repaired", 8, False),
    ("citation_sentence (LLM)", "citation_sentence", 45, False),
    ("citation footnote / reference (which source this row is)", "source_hint", 40, False),
    ("prev_clean", "prev_clean", 35, False), ("core_clean", "core_clean", 45, False),
    ("next_clean", "next_clean", 35, False),
    ("prev_boundary", "prev_boundary", 11, False), ("next_boundary", "next_boundary", 11, False),
    ("source_window (LLM span in « », core in ⟦ ⟧)", "window", 70, False),
    ("match_ok", None, 10, True), ("core_ok", None, 10, True), ("core_complete", None, 10, True),
    ("prev_correct", None, 10, True), ("next_correct", None, 10, True), ("notes", None, 30, True),
    ("citation_id", "citation_id", 38, False), ("comment_id", "comment_id", 22, False),
    ("source_path", "source_path", 40, False),
]
LABELS = {"match_ok": '"Y,Partial,N"', "core_ok": '"Y,N"', "core_complete": '"Y,Cut,Multi,Junk"',
          "prev_correct": '"Y,N"', "next_correct": '"Y,N"'}


def _earlier_ids() -> set[str]:
    ids = set()
    for path in EARLIER:
        if not path.exists():
            continue
        ws = load_workbook(path, read_only=True)["Validation"]
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        i = hdr.index("citation_id")
        ids |= {r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i]}
    return ids


def _sample(con: sqlite3.Connection) -> list[dict]:
    rng, taken_ids, taken_windows, rows = random.Random(SEED), _earlier_ids(), set(), []
    for stratum, n, where in STRATA:
        pool, seen = [], set()
        for r in con.execute(
                "SELECT v.*, c.citation_sentence, c.citation_footnote_text, c.citation_reference_text, "
                "c.citation_link FROM citation_context_v2 v JOIN comment_citations c USING(citation_id) "
                f"WHERE v.core_method IS NOT NULL AND {where} ORDER BY v.citation_id"):
            if r["window_key"] in taken_windows or r["window_key"] in seen or r["citation_id"] in taken_ids:
                continue
            seen.add(r["window_key"])
            pool.append(r)
        for r in rng.sample(pool, min(n, len(pool))):
            taken_ids.add(r["citation_id"])
            taken_windows.add(r["window_key"])
            rows.append({**dict(r), "stratum": stratum})
    return rows


def _source_hint(r: dict) -> str:
    parts = [(lbl, r[k]) for lbl, k in (("footnote", "citation_footnote_text"),
                                        ("reference", "citation_reference_text"),
                                        ("link", "citation_link")) if r[k]]
    return "\n".join(f"[{lbl}] {v[:400]}" for lbl, v in parts) or "(no footnote/reference text)"


def _window(raw: str, r: dict) -> str:
    marks = [(r["span_char_start"], 0, "«"), (r["span_char_end"], 3, "»"),
             (r["core_char_start"], 1, "⟦"), (r["core_char_end"], 2, "⟧")]
    edges = [p for p, _, _ in marks] + [x for x in (r["prev_char_start"], r["next_char_end"]) if x is not None]
    lo, hi = max(0, min(edges) - WINDOW), min(len(raw), max(edges) + WINDOW)
    out, pos = [], lo
    for p, _, m in sorted(marks):
        out += [raw[pos:p], m]
        pos = p
    out.append(raw[pos:hi])
    return (("…" if lo else "") + "".join(out) + ("…" if hi < len(raw) else ""))[:MAX_CELL]


def _write(rows: list[dict]) -> None:
    wb = Workbook()
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill, label_fill = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("Citation context validation set 3 (after the set-2 fixes)", bold),
        (f"{len(rows)} rows from citation_context_v2, one per distinct window, stratified by the new behaviour "
         f"(seed {SEED}). Rows in validation sets 1–2 are excluded. "
         "Regenerate with: python -m validation.make_citation_context_v3_sample", font),
        ("", font),
        ("What changed: every citation now gets ONE core sentence and core ± 1 context. The core is the sentence "
         "carrying the citation's footnote marker when its number is known (from the LLM sentence, the footnote "
         "text, or the footnote the match landed in). relocated = 1 means that sentence lies outside the LLM's "
         "span — the LLM attached the citation to the wrong sentence, or its text matched a footnote. "
         "repaired = 1 means a sentence broken by a page break / footnote block was rejoined.", font),
        ("source_window: the matched LLM span is in « », the chosen core in ⟦ ⟧. prev/core/next are shown "
         "markup-free (the *_clean columns used for hedging).", font),
        ("", font),
        ("Edit only the yellow columns on the 'Validation' tab.", bold),
        ("match_ok — is « » the same text as citation_sentence (LLM)?  Y · Partial (cut short / runs over) · N", font),
        ("core_ok — is ⟦core⟧ the sentence that cites this source (use the footnote/reference column)?  Y · N", font),
        ("core_complete — is core_clean exactly one whole sentence?  Y · Cut (sentence cut off) · "
         "Multi (several sentences) · Junk (footnote, URL, header or table text mixed in)", font),
        ("prev_correct / next_correct — the true neighbouring sentence, or correctly empty at a heading / table / "
         "edge?  Y · N (wrong sentence, footnote/table noise, or empty when real text was adjacent)", font),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ins.cell(i, 1, t).font = f
    ins.column_dimensions["A"].width = 160

    ws = wb.create_sheet("Validation")
    for j, (h, _, w, is_label) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, h)
        c.font, c.alignment, c.border = bold, wrap, thin
        c.fill = label_fill if is_label else head_fill
        ws.column_dimensions[c.column_letter].width = w
    for i, r in enumerate(rows, 2):
        r["val_id"] = i - 1
        for j, (h, k, _, is_label) in enumerate(COLUMNS, 1):
            v = r.get(k) if k else None
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap, thin
            if is_label:
                c.fill = label_fill
    last = len(rows) + 1
    headers = [h for h, *_ in COLUMNS]
    for name, choices in LABELS.items():
        letter = ws.cell(1, headers.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "K2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last}"

    sm = wb.create_sheet("Summary")
    col = {h: ws.cell(1, headers.index(h) + 1).column_letter for h in headers}
    rng = lambda h: f"Validation!${col[h]}$2:${col[h]}${last}"   # noqa: E731
    S = rng("stratum")
    hdr = ["stratum", "n", "labelled", "match Y rate", "core Y rate", "complete Y rate", "prev Y rate", "next Y rate"]
    for j, h in enumerate(hdr, 1):
        c = sm.cell(1, j, h)
        c.font, c.fill = bold, head_fill
        sm.column_dimensions[c.column_letter].width = 16
    rate = lambda crit, h: (f'=IFERROR(COUNTIFS({S},{crit},{rng(h)},"Y")'                  # noqa: E731
                            f'/COUNTIFS({S},{crit},{rng(h)},"<>"),"")')
    for i, s in enumerate([s for s, _, _ in STRATA] + ["ALL"], 2):
        crit = '"*"' if s == "ALL" else f"$A{i}"
        sm.cell(i, 1, s).font = bold if s == "ALL" else font
        sm.cell(i, 2, f"=COUNTIF({S},{crit})")
        sm.cell(i, 3, f'=COUNTIFS({S},{crit},{rng("core_ok")},"<>")')
        for j, h in enumerate(("match_ok", "core_ok", "core_complete", "prev_correct", "next_correct"), 4):
            c = sm.cell(i, j, rate(crit, h))
            c.number_format = "0.0%"
    sm.cell(len(STRATA) + 4, 1, "Rates exclude blank labels; Partial / Cut / Multi / Junk count as not Y.").font = \
        Font(name="Arial", size=9, italic=True)

    wb.save(OUT)


def main() -> None:
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    rows = _sample(con)
    raws: dict[str, str] = {}
    for r in rows:
        p = r["source_path"]
        if p not in raws:
            raws[p] = (MD_ROOT / p).read_text(encoding="utf-8", errors="ignore")
        r["window"] = _window(raws[p], r)
        r["source_hint"] = _source_hint(r)
    _write(rows)
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["stratum"]] = counts.get(r["stratum"], 0) + 1
    print(f"wrote {len(rows)} rows to {OUT}")
    print("  " + "  ".join(f"{k}={v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
