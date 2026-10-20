"""
validation/make_citation_context_sample.py
Draws a stratified, seeded sample of citation_context rows for hand labelling
and writes validation/citation_context_validation.xlsx.

    python -m validation.make_citation_context_sample

Strata are drawn in order without replacement, so a row appears once, under
the first stratum it was drawn for. Each row carries a window of the raw
markdown around the match (match wrapped in ⟦ ⟧) so the labeller can judge
the match and both neighbours without opening the file. Unmatched rows get
the best-scoring fuzzy window (no cutoff) so the labeller can confirm the
sentence really is absent.
"""
import os
import random
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from rapidfuzz import fuzz

from loaders.citation_context import MD_ROOT, Doc, _norm_query

SEED   = 20261012
DB     = Path(__file__).parent.parent / "dockets.db"
OUT    = Path(__file__).parent / "citation_context_validation.xlsx"
WINDOW = 600          # raw chars of markdown shown either side of the match
MAX_CELL = 32000      # Excel cell limit is 32,767

_BASE = "FROM citation_context x JOIN comment_citations c USING(citation_id) WHERE "
# (stratum, n, WHERE clause)
STRATA = [
    ("exact",              30, "x.match_method = 'exact' AND x.n_hits = 1 AND x.unit_type = 'body'"),
    ("exact_body_flow",    15, "x.match_method = 'exact_body_flow'"),
    ("anchor",             15, "x.match_method = 'anchor'"),
    ("fuzzy_95plus",       15, "x.match_method = 'fuzzy' AND x.match_score >= 95"),
    ("fuzzy_85_95",        20, "x.match_method = 'fuzzy' AND x.match_score < 95"),
    ("heading_boundary",   15, "x.source_sentence IS NOT NULL AND x.unit_type = 'body' "
                               "AND (x.prev_boundary = 'heading' OR x.next_boundary = 'heading')"),
    ("multi_hit",          10, "x.source_sentence IS NOT NULL AND x.n_hits > 1"),
    ("cited_heading",       5, "x.unit_type = 'heading'"),
    ("footnote",            5, "x.unit_type = 'footnote'"),
    ("table",               5, "x.unit_type = 'table'"),
    ("unmatched",          20, "x.match_method = 'unmatched'"),
]

COLUMNS = [  # (header, width, is_label)
    ("val_id", 8, False), ("stratum", 16, False), ("match_method", 14, False),
    ("match_score", 9, False), ("n_hits", 7, False), ("unit_type", 10, False),
    ("citation_sentence (LLM)", 45, False),
    ("prev_sentence", 40, False), ("source_sentence", 45, False), ("next_sentence", 40, False),
    ("prev_boundary", 12, False), ("next_boundary", 12, False),
    ("source_window (raw markdown, match in ⟦ ⟧)", 70, False),
    ("match_correct", 12, True), ("prev_correct", 12, True), ("next_correct", 12, True),
    ("truly_absent", 12, True), ("notes", 30, True),
    ("citation_id", 38, False), ("comment_id", 22, False), ("source_path", 40, False),
]
LABELS = {
    "match_correct": '"Y,N,Partial,NA"',
    "prev_correct":  '"Y,N,NA"',
    "next_correct":  '"Y,N,NA"',
    "truly_absent":  '"Y,N,NA"',
}


def _sample(con: sqlite3.Connection) -> list[dict]:
    rng, taken, rows = random.Random(SEED), set(), []
    cols = ("x.citation_id, x.comment_id, x.source_path, x.match_method, x.match_score, x.n_hits, "
            "x.unit_type, x.char_start, x.char_end, x.source_sentence, x.prev_sentence, "
            "x.next_sentence, x.prev_boundary, x.next_boundary, c.citation_sentence")
    for stratum, n, where in STRATA:
        pool = [r for r in con.execute(f"SELECT {cols} {_BASE}{where} ORDER BY x.citation_id")
                if r["citation_id"] not in taken]
        for r in rng.sample(pool, min(n, len(pool))):
            taken.add(r["citation_id"])
            rows.append({**dict(r), "stratum": stratum})
    return rows


def _window(raw: str, a: int, b: int) -> str:
    lo, hi = max(0, a - WINDOW), min(len(raw), b + WINDOW)
    return (("…" if lo else "") + raw[lo:a] + "⟦" + raw[a:b] + "⟧" + raw[b:hi]
            + ("…" if hi < len(raw) else ""))[:MAX_CELL]


def _attach_windows(rows: list[dict]) -> None:
    docs: dict[str, tuple[str, Doc | None]] = {}
    for r in rows:
        if not r["source_path"]:
            r["window"] = None
            continue
        path = MD_ROOT / r["source_path"]
        if r["source_path"] not in docs:
            docs[r["source_path"]] = (path.read_text(encoding="utf-8", errors="ignore"), None)
        raw, doc = docs[r["source_path"]]
        if r["char_start"] is not None:
            r["window"] = _window(raw, r["char_start"], r["char_end"])
            continue
        # unmatched: show the closest fuzzy window, whatever its score
        if doc is None:
            doc = Doc(raw)
            docs[r["source_path"]] = (raw, doc)
        q = _norm_query(r["citation_sentence"] or "")
        al = fuzz.partial_ratio_alignment(q, doc.all_key) if q and doc.all_key else None
        pos = [p for p in doc.all_map[al.dest_start:al.dest_end] if p >= 0] if al else []
        r["window"] = (f"[best fuzzy score {al.score:.0f}] " + _window(raw, min(pos), max(pos) + 1)
                       if pos else "[no candidate window]")


def _write(rows: list[dict]) -> None:
    wb = Workbook()
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_fill  = PatternFill("solid", fgColor="D9D9D9")
    label_fill = PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    # ── Instructions ──
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("Citation context validation set", bold),
        (f"{len(rows)} rows from citation_context, stratified by match method / score band / boundary type "
         f"(seed {SEED}). Regenerate with: python -m validation.make_citation_context_sample", font),
        ("", font),
        ("What to edit: only the yellow columns on the 'Validation' tab (match_correct … notes). "
         "Everything else is generated.", bold),
        ("", font),
        ("match_correct — does source_sentence correspond to the LLM citation_sentence?", font),
        ("    Y = same sentence(s) · Partial = overlaps but too much/too little · N = wrong place · "
         "NA = unmatched row", font),
        ("prev_correct / next_correct — is the neighbour the true preceding/following sentence in the document?", font),
        ("    Y = correct sentence, or correctly NULL (boundary is right: a heading/table/edge really is there)", font),
        ("    N = wrong sentence, a heading leaked in, or NULL when real body text was adjacent · "
         "NA = unmatched row", font),
        ("truly_absent — unmatched rows only: Y if the sentence is genuinely not in the markdown "
         "(check source_window), N if it is there and was missed. NA otherwise.", font),
        ("", font),
        ("Rule being validated: context stops at headings (and tables); headings are never returned as context. "
         "Page-bottom footnotes are skipped when walking body text.", font),
        ("source_window shows ±600 raw markdown chars with the matched span in ⟦ ⟧.", font),
        ("", font),
        ("Example labelled row", bold),
    ]
    for i, (text, f) in enumerate(lines, 1):
        ins.cell(i, 1, text).font = f
    ex_hdr = ["stratum", "citation_sentence (LLM)", "prev_sentence", "source_sentence", "next_sentence",
              "match_correct", "prev_correct", "next_correct", "truly_absent", "notes"]
    ex_row = ["exact", "As noted in our June 2023 comments to NTIA, these systems do not run massive neural networks…",
              "- **Narrow AI applications** have been used safely for decades, can be audited, and pose limited risks.",
              "As noted in our June 2023 comments to NTIA, these systems do not run massive neural networks…",
              "They apply specified or defined rules that have little to no risk associated with them.",
              "Y", "Y", "Y", "NA", "bullet marker kept in prev; fine"]
    r0 = len(lines) + 1
    for j, (h, v) in enumerate(zip(ex_hdr, ex_row), 1):
        ins.cell(r0, j, h).font = bold
        c = ins.cell(r0 + 1, j, v)
        c.font, c.alignment = font, wrap
    ins.column_dimensions["A"].width = 18
    for col in "BCDE":
        ins.column_dimensions[col].width = 40

    # ── Validation ──
    ws = wb.create_sheet("Validation")
    for j, (h, w, is_label) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, h)
        c.font, c.alignment, c.border = bold, wrap, thin
        c.fill = label_fill if is_label else head_fill
        ws.column_dimensions[c.column_letter].width = w
    ws.cell(1, 13).comment = Comment("±600 raw markdown chars around the match; matched span in ⟦ ⟧. "
                                     "Unmatched rows show the best fuzzy candidate and its score.", "generator")
    keys = ["val_id", "stratum", "match_method", "match_score", "n_hits", "unit_type", "citation_sentence",
            "prev_sentence", "source_sentence", "next_sentence", "prev_boundary", "next_boundary", "window",
            None, None, None, None, None, "citation_id", "comment_id", "source_path"]
    for i, r in enumerate(rows, 2):
        r["val_id"] = i - 1
        for j, k in enumerate(keys, 1):
            v = r.get(k) if k else None
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap, thin
            if COLUMNS[j - 1][2]:
                c.fill = label_fill
        if r["match_method"] == "unmatched":   # neighbour labels don't apply
            for j in (14, 15, 16):
                ws.cell(i, j, "NA")
    last = len(rows) + 1
    for name, choices in LABELS.items():
        col = [h for h, _, _ in COLUMNS].index(name) + 1
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        letter = ws.cell(1, col).column_letter
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last}"

    # ── Summary (live formulas over the Validation tab) ──
    sm = wb.create_sheet("Summary")
    hdr = ["stratum", "n", "labelled", "match Y", "match Y rate", "prev Y rate", "next Y rate", "truly absent rate"]
    for j, h in enumerate(hdr, 1):
        c = sm.cell(1, j, h)
        c.font, c.fill = bold, head_fill
        sm.column_dimensions[c.column_letter].width = 16
    rng = lambda col: f"Validation!${col}$2:${col}${last}"   # noqa: E731
    S, M, P, N, A = rng("B"), rng("N"), rng("O"), rng("P"), rng("Q")
    strata = [s for s, _, _ in STRATA] + ["ALL"]
    for i, s in enumerate(strata, 2):
        sm.cell(i, 1, s).font = bold if s == "ALL" else font
        crit = f'"*"' if s == "ALL" else f"$A{i}"
        f = {
            2: f"=COUNTIF({S},{crit})",
            3: f'=COUNTIFS({S},{crit},{M},"<>")',
            4: f'=COUNTIFS({S},{crit},{M},"Y")',
            5: f'=IFERROR(D{i}/(COUNTIFS({S},{crit},{M},"<>")-COUNTIFS({S},{crit},{M},"NA")),"")',
            6: f'=IFERROR(COUNTIFS({S},{crit},{P},"Y")/(COUNTIFS({S},{crit},{P},"<>")-COUNTIFS({S},{crit},{P},"NA")),"")',
            7: f'=IFERROR(COUNTIFS({S},{crit},{N},"Y")/(COUNTIFS({S},{crit},{N},"<>")-COUNTIFS({S},{crit},{N},"NA")),"")',
            8: f'=IFERROR(COUNTIFS({S},{crit},{A},"Y")/(COUNTIFS({S},{crit},{A},"<>")-COUNTIFS({S},{crit},{A},"NA")),"")',
        }
        for j, formula in f.items():
            c = sm.cell(i, j, formula)
            c.font = font
            if j >= 5:
                c.number_format = "0.0%"
    n = len(strata) + 3
    sm.cell(n, 1, "Rates exclude blank and NA labels. Rows are strata from the generator (see Instructions); "
                  "'match Y rate' treats Partial as not correct.").font = Font(name="Arial", size=9, italic=True)

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)


def main() -> None:
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    rows = _sample(con)
    _attach_windows(rows)
    _write(rows)
    counts = {}
    for r in rows:
        counts[r["stratum"]] = counts.get(r["stratum"], 0) + 1
    print(f"wrote {len(rows)} rows to {OUT}")
    print("  " + "  ".join(f"{k}={v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
