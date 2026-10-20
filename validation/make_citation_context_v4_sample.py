"""
validation/make_citation_context_v4_sample.py
Draws a stratified, seeded sample of citation_context_v3 rows for hand
labelling and writes validation/citation_context_v4_validation.xlsx.

    python -m validation.make_citation_context_v4_sample

Strata target the v3 changes: footnote citations recovered to the body
sentence the LLM missed, duplicates of sentences the LLM already captured,
the new reference / table / footnote / bare-link exclusions (and in-scope
rows that look reference-like, for recall), cores fixed by marker-safe
splitting, and a regression check on low-confidence in-scope rows.
Citations in validation sets 1–3 are skipped; one citation per window_key.
"""
import random
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from loaders.citation_context import MD_ROOT
from loaders.citation_context_v3 import REF_LINE_CUES, ref_cues
from validation.make_citation_context_v3_sample import _source_hint, _window

SEED   = 20261015
DB     = Path(__file__).parent.parent / "dockets.db"
OUT    = Path(__file__).parent / "citation_context_v4_validation.xlsx"
EARLIER = [Path(__file__).parent / f for f in ("citation_context_validation.xlsx",
                                               "citation_context_v2_validation.xlsx",
                                               "citation_context_v3_validation.xlsx")]
MAX_CELL = 32000

_V2_SPLIT = ("(o.core_clean GLOB '[0-9]*]*' OR o.core_clean GLOB '[0-9][0-9]*, [0-9]*]*' "
             "OR o.prev_clean LIKE '% et al.' OR o.core_clean GLOB '[ivx]* [A-Z]*')")
# (stratum, n, WHERE clause on v = citation_context_v3 / o = citation_context_v2, python filter or None)
STRATA = [
    ("recovered",          25, "v.analysis_scope = 'recovered'", None),
    ("excluded_duplicate", 15, "v.analysis_scope = 'excluded_duplicate'", None),
    ("excluded_reference", 20, "v.analysis_scope = 'excluded_reference'", None),
    ("in_scope_ref_like",  10, "v.analysis_scope = 'in_scope'",
     lambda r: ref_cues(r["citation_sentence"] or "") >= REF_LINE_CUES),
    ("excluded_table",      5, "v.analysis_scope = 'excluded_table'", None),
    ("excluded_footnote",   5, "v.analysis_scope = 'excluded_footnote'", None),
    ("excluded_bare_link",  5, "v.analysis_scope = 'excluded_bare_link'", None),
    ("marker_split_fixed", 15, f"v.analysis_scope = 'in_scope' AND {_V2_SPLIT}", None),
    ("in_scope_low_conf",  20, "v.analysis_scope = 'in_scope' AND (v.core_method = 'span_last' "
                               "OR (v.match_method = 'fuzzy' AND v.match_score < 90))", None),
]
EXCLUDED = {s for s, *_ in STRATA if s.startswith("excluded_") and s != "excluded_duplicate"}

COLUMNS = [  # (header, key, width, is_label)
    ("val_id", "val_id", 7, False), ("stratum", "stratum", 16, False),
    ("analysis_scope", "analysis_scope", 16, False), ("match_unit_type", "match_unit_type", 10, False),
    ("match_method", "match_method", 11, False), ("match_score", "match_score", 7, False),
    ("core_method", "core_method", 14, False), ("citation_number", "citation_number", 8, False),
    ("citation_sentence (LLM)", "citation_sentence", 45, False),
    ("citation footnote / reference (which source this row is)", "source_hint", 40, False),
    ("prev_clean", "prev_clean", 35, False), ("core_clean", "core_clean", 45, False),
    ("next_clean", "next_clean", 35, False),
    ("v2 core_clean (before)", "v2_core", 35, False),
    ("duplicate_of: LLM sentence already covering the core", "dup_sentence", 40, False),
    ("source_window (LLM span in « », core in ⟦ ⟧)", "window", 70, False),
    ("scope_correct", None, 10, True), ("core_ok", None, 10, True), ("core_complete", None, 10, True),
    ("prev_correct", None, 10, True), ("next_correct", None, 10, True), ("notes", None, 30, True),
    ("citation_id", "citation_id", 38, False), ("comment_id", "comment_id", 22, False),
    ("source_path", "source_path", 40, False),
]
LABELS = {"scope_correct": '"Y,N"', "core_ok": '"Y,N,NA"', "core_complete": '"Y,Cut,Multi,Junk,NA"',
          "prev_correct": '"Y,N,NA"', "next_correct": '"Y,N,NA"'}
EXAMPLE = {"stratum": "recovered", "scope_correct": "Y", "core_ok": "Y", "core_complete": "Y",
           "prev_correct": "Y", "next_correct": "N", "notes": "next runs into a page footer"}


def _earlier_ids() -> set[str]:
    ids = set()
    for path in EARLIER:
        if not path.exists():
            continue
        ws = load_workbook(path, read_only=True)["Validation"]
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        i = hdr.index("citation_id")
        ids |= {r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i]}
    return ids


def _sample(con: sqlite3.Connection) -> list[dict]:
    rng, taken_ids, taken_windows, rows = random.Random(SEED), _earlier_ids(), set(), []
    for stratum, n, where, keep in STRATA:
        pool, seen = [], set()
        for r in con.execute(
                "SELECT v.*, o.core_clean AS v2_core, c.citation_sentence, c.citation_footnote_text, "
                "c.citation_reference_text, c.citation_link, d.citation_sentence AS dup_sentence "
                "FROM citation_context_v3 v JOIN comment_citations c USING(citation_id) "
                "LEFT JOIN citation_context_v2 o USING(citation_id) "
                "LEFT JOIN comment_citations d ON d.citation_id = v.duplicate_of "
                f"WHERE v.core_method IS NOT NULL AND {where} ORDER BY v.citation_id"):
            if r["window_key"] in taken_windows or r["window_key"] in seen or r["citation_id"] in taken_ids:
                continue
            if keep is not None and not keep(r):
                continue
            seen.add(r["window_key"])
            pool.append(r)
        for r in rng.sample(pool, min(n, len(pool))):
            taken_ids.add(r["citation_id"])
            taken_windows.add(r["window_key"])
            rows.append({**dict(r), "stratum": stratum})
    return rows


def _write(rows: list[dict]) -> None:
    wb = Workbook()
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill, label_fill = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("Citation context validation set 4 (v3: analysis scope)", bold),
        (f"{len(rows)} rows from citation_context_v3, one per distinct window (seed {SEED}). Rows in validation "
         "sets 1–3 are excluded. Regenerate with: python -m validation.make_citation_context_v4_sample", font),
        ("", font),
        ("What changed: every citation now has an analysis_scope. Only in_scope (LLM sentence is body text) and "
         "recovered rows are analysed. recovered = the LLM extracted footnote/reference text, and the body "
         "sentence carrying that footnote's marker was NOT captured by any other LLM citation in the comment. "
         "excluded_duplicate = same, but another LLM citation already covers that body sentence. Table, "
         "footnote, reference-list, heading and bare-link text is excluded.", font),
        ("source_window: the matched LLM span is in « », the chosen core in ⟦ ⟧. 'v2 core_clean (before)' shows "
         "the previous version's core for comparison.", font),
        ("", font),
        ("Edit only the yellow columns on the 'Validation' tab. Excluded strata (excluded_reference / table / "
         "footnote / bare_link) need scope_correct only; their other label columns are prefilled NA.", bold),
        ("scope_correct — is analysis_scope right?  Y · N", font),
        ("   recovered: ⟦core⟧ cites this source AND the LLM did not capture that sentence elsewhere", font),
        ("   excluded_duplicate: the 'duplicate_of' LLM sentence is the same sentence as ⟦core⟧", font),
        ("   excluded_*: the LLM sentence really is reference / table / footnote / bare-link text", font),
        ("   in_scope_*: the LLM sentence really is body prose (not a reference entry)", font),
        ("core_ok — is ⟦core⟧ the sentence that cites this source?  Y · N · NA", font),
        ("core_complete — is core_clean exactly one whole sentence?  Y · Cut · Multi · Junk · NA", font),
        ("prev_correct / next_correct — the true neighbouring sentence, or correctly empty at a heading / table / "
         "edge?  Y · N · NA", font),
        ("", font),
        ("Example of a labelled row (format only):", bold),
        ("   " + " · ".join(f"{k}={v}" for k, v in EXAMPLE.items()), font),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ins.cell(i, 1, t).font = f
    ins.column_dimensions["A"].width = 160

    ws = wb.create_sheet("Validation")
    for j, (h, _, w, is_label) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, h)
        c.font, c.alignment, c.border = bold, wrap, thin
        c.fill = label_fill if is_label else head_fill
        ws.column_dimensions[c.column_letter].width = w
    for i, r in enumerate(rows, 2):
        r["val_id"] = i - 1
        for j, (h, k, _, is_label) in enumerate(COLUMNS, 1):
            v = r.get(k) if k else None
            if is_label and r["stratum"] in EXCLUDED and h in ("core_ok", "core_complete",
                                                               "prev_correct", "next_correct"):
                v = "NA"
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap, thin
            if is_label:
                c.fill = label_fill
    last = len(rows) + 1
    headers = [h for h, *_ in COLUMNS]
    for name, choices in LABELS.items():
        letter = ws.cell(1, headers.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "J2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last}"

    sm = wb.create_sheet("Summary")
    col = {h: ws.cell(1, headers.index(h) + 1).column_letter for h in headers}
    rng = lambda h: f"Validation!${col[h]}$2:${col[h]}${last}"   # noqa: E731
    S = rng("stratum")
    hdr = ["stratum", "n", "labelled", "scope Y rate", "core Y rate", "complete Y rate", "prev Y rate", "next Y rate"]
    for j, h in enumerate(hdr, 1):
        c = sm.cell(1, j, h)
        c.font, c.fill = bold, head_fill
        sm.column_dimensions[c.column_letter].width = 18
    # rates ignore blank and NA labels
    rate = lambda crit, h: (f'=IFERROR(COUNTIFS({S},{crit},{rng(h)},"Y")'                  # noqa: E731
                            f'/(COUNTIFS({S},{crit},{rng(h)},"<>")-COUNTIFS({S},{crit},{rng(h)},"NA")),"")')
    names = [s for s, *_ in STRATA]
    for i, s in enumerate(names + ["ALL"], 2):
        crit = '"*"' if s == "ALL" else f"$A{i}"
        sm.cell(i, 1, s).font = bold if s == "ALL" else font
        sm.cell(i, 2, f"=COUNTIF({S},{crit})")
        sm.cell(i, 3, f'=COUNTIFS({S},{crit},{rng("scope_correct")},"<>")')
        for j, h in enumerate(("scope_correct", "core_ok", "core_complete", "prev_correct", "next_correct"), 4):
            c = sm.cell(i, j, rate(crit, h))
            c.number_format = "0.0%"
    sm.cell(len(names) + 4, 1, "Rates exclude blank and NA labels; Cut / Multi / Junk count as not Y.").font = \
        Font(name="Arial", size=9, italic=True)

    wb.save(OUT)


def main() -> None:
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    rows = _sample(con)
    raws: dict[str, str] = {}
    for r in rows:
        p = r["source_path"]
        if p not in raws:
            raws[p] = (MD_ROOT / p).read_text(encoding="utf-8", errors="ignore")
        r["window"] = _window(raws[p], r)
        r["source_hint"] = _source_hint(r)
    _write(rows)
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["stratum"]] = counts.get(r["stratum"], 0) + 1
    print(f"wrote {len(rows)} rows to {OUT}")
    print("  " + "  ".join(f"{k}={v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
