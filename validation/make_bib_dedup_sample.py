"""
validation/make_bib_dedup_sample.py
Draws a stratified, seeded sample of confident bibliography duplicates and
writes validation/bib_dedup_validation.xlsx for hand labelling.

    venv/bin/python -m validation.make_bib_dedup_sample

The unmatched_bibliography rows (citation_context_v2.match_method =
'empty_sentence') were loaded without checking them against the cited
references of the same comment. A row is a confident duplicate when, within
its comment, a cited row (one with a citation sentence) shares
  - a URL with a real path            → url_* strata, or
  - near-identical reference text     → text_* strata
    (same first surname, same year where both give one; token_set ≥ 90 when
    one side is a short form like "Engstrom and Ho 2020", token_sort ≥ 90
    when both are full references).
Borderline "possible" matches (score 80–90) are not sampled. At most
PER_COMMENT rows per comment per stratum.
"""
import random
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from rapidfuzz import fuzz

SEED        = 20261015
DB          = Path(__file__).parent.parent / "dockets.db"
OUT         = Path(__file__).parent / "bib_dedup_validation.xlsx"
PER_COMMENT = 2
MAX_CELL    = 32000

# (stratum, n, description)
STRATA = [
    ("url_full_ref",    20, "full bibliography entry; shares a URL with a cited reference"),
    ("url_bare",        10, "bare URL entry; same URL as a cited reference"),
    ("text_short_form", 15, "text match where one side is a short form (author–year / short title)"),
    ("text_full",       15, "text match between two full references"),
]

COLUMNS = [  # (header, key, width, is_label)
    ("val_id", "val_id", 7, False), ("stratum", "stratum", 15, False),
    ("score", "score", 7, False), ("shared_url", "shared_url", 30, False),
    ("bibliography entry (empty_sentence row)", "bib_text", 50, False),
    ("bib link", "bib_link", 25, False),
    ("matched: citation_sentence (LLM)", "m_sentence", 45, False),
    ("matched: reference / footnote", "m_reference", 45, False),
    ("matched link", "m_link", 25, False),
    ("same_work", None, 10, True), ("pairing_ok", None, 10, True), ("notes", None, 30, True),
    ("bib_citation_id", "citation_id", 38, False), ("matched_citation_id", "m_citation_id", 38, False),
    ("comment_id", "comment_id", 22, False),
]
LABELS = {"same_work": '"Y,N,Unsure"', "pairing_ok": '"Y,N,NA"'}

# ── matching ────────────────────────────────────────────────────────────────
URLS   = re.compile(r"(?:https?://|www\.)[^\s<>\]\)]+")
MDLINK = re.compile(r"\[([^\]]*)\]\((?:[^()\s]|\([^)]*\))*\)")
YR     = re.compile(r"\b(?:19|20)\d{2}\b")
BOILER = set("accessed retrieved available at from url online last visited n d january february march april "
             "may june july august september october november december jan feb mar apr jun jul aug sep sept "
             "oct nov dec pdf html et al".split())
STOP   = {"the", "and", "for", "with", "from", "see", "of", "in", "on", "a", "an", "et", "al"}


def nurl(u: str) -> str | None:
    u = re.split(r"\]\(|\)\[|\s", u.strip().strip("<>[]()"))[0].lower()
    u = re.sub(r"^https?://|^www\.|[?#].*$|[\\/.,;)\]>]+$", "", u).replace("\\", "")
    return u if len(u) > 8 else None


def urls(*texts: str | None) -> set[str]:
    found = set()
    for t in texts:
        for u in URLS.findall(t or "") + ([t] if t and URLS.match(t.strip("<[ ")) else []):
            if (n := nurl(u)) and "/" in n:          # path required: bare homepages are too generic
                found.add(n)
    return found


def ntext(t: str | None) -> str:
    t = URLS.sub(" ", t or "")
    t = re.sub(r"^\s*(?:[-*•]|\[[^\]]{1,20}\]|\d{1,3}[.)]?)\s+", "", t)
    return " ".join(w for w in re.sub(r"[^\w]+", " ", t.lower()).split() if w not in BOILER)


def _surname(t: str) -> str | None:
    return next((w for w in t.split() if w.isalpha() and len(w) > 2 and w not in STOP), None)


def text_score(a: str, b: str) -> tuple[float, bool]:
    """(score, short_form) for two normalised references; 0 when implausible."""
    if len(a) < 15 or len(b) < 15:
        return 0, False
    ya, yb = set(YR.findall(a)), set(YR.findall(b))
    sa, sb = _surname(a), _surname(b)
    if (ya and yb and not ya & yb) or not (sa and sb and (sa in b.split() or sb in a.split())):
        return 0, False
    if min(len(a.split()), len(b.split())) <= 6:
        return (fuzz.token_set_ratio(a, b) if ya and yb else 0), True
    return fuzz.token_sort_ratio(a, b), False


def is_bare(text: str) -> bool:
    t = MDLINK.sub(lambda m: "" if re.match(r"\s*(https?://|www\.)|\S*[/?=&%_]\S*$", m.group(1)) else m.group(1), text)
    t = URLS.sub(" ", t)
    t = re.sub(r"^\s*(?:[-*•]|\[\w{1,20}\]|\d{1,3}[.)]?)\s*", " ", t)
    return len(re.sub(r"[^\w]+", " ", t).split()) <= 2


def candidates(con: sqlite3.Connection) -> dict[str, list[dict]]:
    rows = [dict(r) for r in con.execute(
        """SELECT c.citation_id, c.comment_id, v.match_method, c.citation_sentence,
                  c.unmatched_citation_reference_text, c.citation_reference_text,
                  c.citation_footnote_text, c.citation_link
             FROM citation_context_v2 v JOIN comment_citations c USING (citation_id)""")]
    cited = defaultdict(list)
    for r in rows:
        if r["match_method"] != "empty_sentence":
            cited[r["comment_id"]].append(r)
    pools: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        if r["match_method"] != "empty_sentence":
            continue
        bib = r["unmatched_citation_reference_text"] or r["citation_sentence"] or ""
        b_urls, b_txt = urls(bib, r["citation_link"]), ntext(bib)
        hit = None
        for o in cited[r["comment_id"]]:
            shared = b_urls & urls(o["citation_reference_text"], o["citation_footnote_text"], o["citation_link"])
            if shared:
                hit = ("url_bare" if is_bare(bib) else "url_full_ref", 100.0, o, sorted(shared)[0])
                break
            for field in (o["citation_reference_text"], o["citation_footnote_text"]):
                s, short = text_score(b_txt, ntext(field))
                if s >= 90 and (hit is None or s > hit[1]):
                    hit = ("text_short_form" if short else "text_full", s, o, "")
        if hit:
            stratum, score, o, shared = hit
            pools[stratum].append({
                "stratum": stratum, "score": round(score, 1), "shared_url": shared,
                "bib_text": bib, "bib_link": r["citation_link"], "citation_id": r["citation_id"],
                "comment_id": r["comment_id"], "m_citation_id": o["citation_id"],
                "m_sentence": o["citation_sentence"],
                "m_reference": "\n".join(f"[{k}] {o[f]}" for k, f in (("reference", "citation_reference_text"),
                                                                     ("footnote", "citation_footnote_text"))
                                         if o[f]),
                "m_link": o["citation_link"],
            })
    return pools


def _sample(pools: dict[str, list[dict]]) -> list[dict]:
    rng, out = random.Random(SEED), []
    for stratum, n, _ in STRATA:
        pool = sorted(pools[stratum], key=lambda r: r["citation_id"])
        rng.shuffle(pool)
        per: dict[str, int] = defaultdict(int)
        picked = []
        for r in pool:
            if per[r["comment_id"]] < PER_COMMENT:
                per[r["comment_id"]] += 1
                picked.append(r)
            if len(picked) == n:
                break
        out += picked
    return out


# ── workbook ────────────────────────────────────────────────────────────────
def _write(rows: list[dict], pools: dict[str, list[dict]]) -> None:
    wb = Workbook()
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill, label_fill = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    ins = wb.active
    ins.title = "Instructions"
    total = sum(len(p) for p in pools.values())
    lines = [
        ("Bibliography de-duplication — validation set", bold),
        (f"{len(rows)} of {total} confident duplicates (seed {SEED}; ≤{PER_COMMENT} per comment per stratum).", font),
        ("Each row pairs an unmatched_bibliography entry (no citation sentence) with the cited row of the same "
         "comment it was matched to. If the match is right, the bibliography row is redundant and can be dropped "
         "(after copying its fuller reference text onto the cited row).", font),
        ("", font),
        ("Strata (sampled / pool):", bold),
        *[(f"   {s}: {sum(r['stratum'] == s for r in rows)} / {len(pools[s])} — {d}", font) for s, _, d in STRATA],
        ("", font),
        ("Labels (yellow columns):", bold),
        ("   same_work — do the bibliography entry and the matched reference point to the same work "
         "(same paper / report / page; a different edition, version or chapter counts as N)?  Y · N · Unsure", font),
        ("   pairing_ok — does the matched citation_sentence actually cite that work? N means the cited row is itself "
         "mis-paired, so dropping the bibliography row would lose the work.  Y · N · NA (no sentence to judge)", font),
        ("", font),
        ("Regenerate with: venv/bin/python -m validation.make_bib_dedup_sample", font),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ins.cell(i, 1, t).font = f
    ins.column_dimensions["A"].width = 160

    ws = wb.create_sheet("Validation")
    for j, (h, _, w, is_label) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, h)
        c.font, c.alignment, c.border = bold, wrap, thin
        c.fill = label_fill if is_label else head_fill
        ws.column_dimensions[c.column_letter].width = w
    for i, r in enumerate(rows, 2):
        r["val_id"] = i - 1
        for j, (_, k, _, is_label) in enumerate(COLUMNS, 1):
            v = r.get(k) if k else None
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap, thin
            if is_label:
                c.fill = label_fill
    last = len(rows) + 1
    headers = [h for h, *_ in COLUMNS]
    for name, choices in LABELS.items():
        letter = ws.cell(1, headers.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last}"

    sm = wb.create_sheet("Summary")
    col = {h: ws.cell(1, headers.index(h) + 1).column_letter for h in headers}
    rng = lambda h: f"Validation!${col[h]}$2:${col[h]}${last}"   # noqa: E731
    S = rng("stratum")
    hdr = ["stratum", "pool", "n", "labelled", "same_work Y rate", "pairing_ok Y rate"]
    for j, h in enumerate(hdr, 1):
        c = sm.cell(1, j, h)
        c.font, c.fill = bold, head_fill
        sm.column_dimensions[c.column_letter].width = 18
    rate = lambda crit, h, excl: (f'=IFERROR(COUNTIFS({S},{crit},{rng(h)},"Y")'           # noqa: E731
                                  f'/(COUNTIFS({S},{crit},{rng(h)},"<>")-COUNTIFS({S},{crit},{rng(h)},"{excl}")),"")')
    names = [s for s, *_ in STRATA]
    for i, s in enumerate(names + ["ALL"], 2):
        crit = '"*"' if s == "ALL" else f"$A{i}"
        sm.cell(i, 1, s).font = bold if s == "ALL" else font
        sm.cell(i, 2, total if s == "ALL" else len(pools[s]))
        sm.cell(i, 3, f"=COUNTIF({S},{crit})")
        sm.cell(i, 4, f'=COUNTIFS({S},{crit},{rng("same_work")},"<>")')
        sm.cell(i, 5, rate(crit, "same_work", "Unsure")).number_format = "0.0%"
        sm.cell(i, 6, rate(crit, "pairing_ok", "NA")).number_format = "0.0%"
    sm.cell(len(names) + 4, 1, "Rates exclude blank labels and Unsure / NA. "
            "ALL is an unweighted sample rate; weight by pool for the population estimate.").font = \
        Font(name="Arial", size=9, italic=True)

    wb.save(OUT)


def main() -> None:
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    pools = candidates(con)
    rows = _sample(pools)
    _write(rows, pools)
    print(f"wrote {len(rows)} rows to {OUT}")
    print("  pools:   " + "  ".join(f"{s}={len(pools[s])}" for s, *_ in STRATA))
    print("  sampled: " + "  ".join(f"{s}={sum(r['stratum'] == s for r in rows)}" for s, *_ in STRATA))


if __name__ == "__main__":
    main()
